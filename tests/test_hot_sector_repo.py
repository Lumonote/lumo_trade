from __future__ import annotations

import sqlite3

import pytest

from data_store.schema import migrate


@pytest.fixture
def conn(tmp_path, monkeypatch):
    path = tmp_path / "kronos_test.sqlite"
    c = sqlite3.connect(path, isolation_level=None)
    c.row_factory = sqlite3.Row
    migrate(c)
    _getter = lambda: c  # noqa: E731
    from data_store import hot_sector_repo
    monkeypatch.setattr(hot_sector_repo, "get_conn", _getter)
    yield c
    c.close()


def test_hot_sector_snapshot_roundtrip_and_relations(conn):
    from data_store import hot_sector_repo as repo

    sid = repo.save_snapshot(
        boards=[
            {"code": "BK1001", "name": "半导体", "type": "行业", "rank": 1, "change_pct": 3.2, "main_net_inflow": 1.2e8},
        ],
        stocks=[
            {
                "sector_code": "BK1001", "code": "688001", "name": "芯片测试",
                "sector_stock_rank": 1, "candidate_rank": 1, "price": 20.0,
                "change_pct": 4.0, "main_net_inflow": 3e6,
                "main_net_inflow_text": "300.0万", "lhb_trade_date": "2026-06-13",
                "lhb_buy_amount": 1.2e8, "lhb_net_amount": 6e7,
                "lhb_reason": "日涨幅偏离值达7%",
            },
        ],
        relations=[
            {
                "board_code": "BK1001", "code": "688001", "relation_type": "dragon_tiger",
                "related_table": "dragon_tiger_list", "related_key": "2026-06-13:688001",
                "trade_date": "2026-06-13", "amount": 6e7,
            },
        ],
        meta={"source": "sector_hot", "board_limit": 10},
    )

    latest = repo.latest_snapshot()
    assert latest["id"] == sid
    assert latest["stock_count"] == 1
    boards = repo.boards_for_snapshot(sid)
    assert boards[0]["board_name"] == "半导体"
    assert boards[0]["relation_count"] == 1
    stocks = repo.stocks_for_snapshot(sid, board_code="BK1001")
    assert stocks[0]["code"] == "688001"
    assert stocks[0]["lhb_buy_amount"] == 1.2e8
    relations = repo.relations_for_snapshot(sid)
    assert relations[0]["relation_type"] == "dragon_tiger"


def test_hot_sector_snapshot_history_order(conn):
    from data_store import hot_sector_repo as repo

    older = repo.save_snapshot(
        boards=[{"code": "BK1001", "name": "半导体", "rank": 1}],
        stocks=[],
        relations=[],
        meta={"created_at": "2026-06-13T09:30:00", "source": "sector_hot"},
    )
    newer = repo.save_snapshot(
        boards=[{"code": "BK2001", "name": "机器人", "rank": 1}],
        stocks=[],
        relations=[],
        meta={"created_at": "2026-06-14T09:30:00", "source": "sector_hot"},
    )

    rows = repo.list_snapshots(limit=5)

    assert [row["id"] for row in rows[:2]] == [newer, older]


def test_hot_sector_export_excel_has_dimension_sheets(conn, tmp_path):
    from data_store import hot_sector_repo as repo

    sid = repo.save_snapshot(
        boards=[{"code": "BK1001", "name": "半导体", "type": "行业", "rank": 1}],
        stocks=[{"sector_code": "BK1001", "code": "688001", "name": "芯片测试", "sector_stock_rank": 1}],
        relations=[],
        meta={"source": "sector_hot"},
    )
    out = repo.export_snapshot_excel(sid, tmp_path / "hot_sector.xlsx")

    assert out.exists()
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    assert {"快照", "热门板块", "成分股排名资金", "龙虎榜命中", "关联关系", "半导体"} <= set(wb.sheetnames)
    wb.close()


def test_hot_sector_export_excel_handles_missing_board_rows(conn, tmp_path):
    from data_store import hot_sector_repo as repo

    sid = repo.save_snapshot(
        boards=[],
        stocks=[{"sector_code": "BK9999", "code": "688001", "name": "芯片测试", "sector_stock_rank": 1}],
        relations=[],
        meta={"source": "sector_hot"},
    )
    out = repo.export_snapshot_excel(sid, tmp_path / "hot_sector_missing_board.xlsx")

    assert out.exists()
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    assert "BK9999" in wb.sheetnames
    wb.close()
