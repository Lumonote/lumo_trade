import importlib
import sqlite3
import sys


def _load_webui_core(tmp_path, monkeypatch):
    monkeypatch.setenv("KRONOS_USER_DIR", str(tmp_path))
    # configuration_service does os.environ.setdefault("KRONOS_RESULTS_DIR", ...)
    # on first import, which pins results_dir() to the first tmp dir and breaks
    # re-import isolation across tests. Align it explicitly with this tmp dir.
    monkeypatch.setenv("KRONOS_RESULTS_DIR", str(tmp_path / "results"))
    monkeypatch.setenv("KRONOS_DISABLE_TORCH", "1")
    sys.modules.pop("webui.core", None)
    module = importlib.import_module("webui.core")
    return module


def _patch_store_conn(monkeypatch, conn):
    from data_store import hot_sector_repo, opportunity_repo
    monkeypatch.setattr(hot_sector_repo, "get_conn", lambda: conn)
    monkeypatch.setattr(opportunity_repo, "get_conn", lambda: conn)


def test_latest_opportunity_report_parser_handles_generated_html_table(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = results_dir / "opportunity_top10_20260521_101500.md"
    report.write_text(
        """
## 风险提示
<p>⚠️ <strong>今日候选质量偏低</strong></p>

## 综合排名 TOP20
<table>
<thead><tr><th>排名</th><th>代码</th><th>股票名称</th><th>综合得分</th><th>详细分析</th></tr></thead>
<tbody>
<tr><td>1</td><td>603005</td><td>晶方科技</td><td>83.25</td><td>【概览】评级A，建议：可关注；【板块】半导体(+3.2%, 强势领涨, 82分)；【量化】买2/卖0/总2(—)，75分，模型[RSI,MACD]；【技术】RSI:55，MACD:金叉；【情绪资金】情绪:中性；【关键加减分】量价改善；【入选原因】半导体板块走强</td></tr>
</tbody></table>
""",
        encoding="utf-8",
    )

    payload = module._load_latest_opportunities()

    assert payload["latest_report"]["file"] == report.name
    assert payload["market_env"] == "今日候选质量偏低"
    assert payload["stats"]["total"] == 1
    assert payload["stats"]["strong_count"] == 1
    assert payload["items"][0]["code"] == "603005"
    assert payload["items"][0]["sector"] == "半导体"
    assert payload["items"][0]["quant_models"] == ["RSI", "MACD"]
    assert [item["name"] for item in payload["quant_models"]] == ["RSI", "MACD"]


_RICH_DETAIL = (
    "【概览】评级S，建议：🌟 强烈推荐；"
    "【涨幅】当日:+7.38%，3日:+8.28%，5日:+4.39%；"
    "【板块】软件服务(—, 市场震荡, 50分)；"
    "【量化】买6/卖1/总30(20%)，97分，模型[机器学习RF, 资金趋势]；"
    "【技术】RSI:45.79，MACD:死叉，布林:中轨上方，54分；"
    "【基本面】营收:+23.9%，利润:+445.0%，35分；"
    "【情绪资金】情绪:中性，60分；"
    "【消息】评级:中性，利好0/利空0，50分；"
    "【关键加减分】牛股动量识别: 放量上涨2日:+5 (总+5分)；"
    "【入选原因】综合评级S级(强烈推荐)；"
    "【最新动态】金山办公软件激励计划自查报告；"
    "【高级】追高风险: 🟢 低(0分) 高级评分: 56.9分"
)


def _write_rich_report(results_dir, name="opportunity_top10_20260604_212907.md"):
    report = results_dir / name
    report.write_text(
        "## 综合排名 TOP20\n"
        "<table>\n"
        "<thead><tr><th>排名</th><th>代码</th><th>股票名称</th><th>综合得分</th><th>详细分析</th></tr></thead>\n"
        "<tbody>\n"
        f"<tr><td>3</td><td>688111</td><td>金山办公</td><td>87.08</td><td>{_RICH_DETAIL}</td></tr>\n"
        "</tbody></table>\n",
        encoding="utf-8",
    )
    return report


def test_extract_all_detail_sections_returns_ordered_label_value(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)

    sections = module._extract_all_detail_sections(_RICH_DETAIL)

    labels = [section["label"] for section in sections]
    assert labels == [
        "概览", "涨幅", "板块", "量化", "技术", "基本面",
        "情绪资金", "消息", "关键加减分", "入选原因", "最新动态", "高级",
    ]
    by_label = {section["label"]: section["value"] for section in sections}
    assert by_label["涨幅"] == "当日:+7.38%，3日:+8.28%，5日:+4.39%"
    assert by_label["概览"].startswith("评级S")
    # value must be trimmed of the trailing section separator
    assert not by_label["板块"].endswith("；")


def test_extract_all_detail_sections_folds_embedded_markers(tmp_path, monkeypatch):
    """新闻/动态正文里内嵌的 【…】 不能变成顶层字段(否则污染画布与 Excel 列)。"""
    module = _load_webui_core(tmp_path, monkeypatch)

    detail = (
        "【概览】评级A，建议：可关注；"
        "【最新动态】公司公告 【股商异动】今日涨停，主力净流入1.6亿；"
        "【消息】评级:中性，利好1/利空0，58分"
    )
    sections = module._extract_all_detail_sections(detail)

    labels = [section["label"] for section in sections]
    assert labels == ["概览", "最新动态", "消息"]
    # embedded 【股商异动】 stays inside 最新动态's value, not a separate field
    assert "股商异动" not in labels
    by_label = {section["label"]: section["value"] for section in sections}
    assert "股商异动" in by_label["最新动态"]
    assert "主力净流入1.6亿" in by_label["最新动态"]


def test_parse_opportunity_report_items_include_full_fields(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = _write_rich_report(results_dir)

    parsed = module._parse_opportunity_report(report)

    item = parsed["items"][0]
    assert item["code"] == "688111"
    labels = [field["label"] for field in item["fields"]]
    assert "涨幅" in labels and "板块" in labels and "高级" in labels
    # existing keys must remain intact (no regression for current callers)
    assert item["sector"] == "软件服务"
    assert item["quant_models"] == ["机器学习RF", "资金趋势"]


def test_load_opportunity_report_cards_by_file(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = _write_rich_report(results_dir)

    payload = module.load_opportunity_report_cards(report.name)

    assert payload is not None
    assert payload["file"] == report.name
    assert payload["count"] == 1
    card = payload["cards"][0]
    assert card["code"] == "688111"
    assert card["score"] == 87.08
    assert card["rating"] == "S"
    assert any(field["label"] == "板块" for field in card["fields"])


def test_load_opportunity_report_cards_rejects_traversal_and_missing(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    _write_rich_report(results_dir)

    assert module.load_opportunity_report_cards("../../etc/passwd") is None
    assert module.load_opportunity_report_cards("does_not_exist.md") is None
    assert module.load_opportunity_report_cards("") is None
    # non-opportunity markdown in the results dir is not served by this endpoint
    (results_dir / "random_notes.md").write_text("hello", encoding="utf-8")
    assert module.load_opportunity_report_cards("random_notes.md") is None


def test_parse_opportunity_report_excludes_review_and_confidence_tables(tmp_path, monkeypatch):
    """Real reports start with a 昨日复盘 table (8 cols, detail = buy price, no 【】).

    Those rows must NOT be treated as opportunity items — only the TOP20 ranking
    rows (whose detail cell carries 【概览】…) are real candidates.
    """
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = results_dir / "opportunity_top10_20260602_095017.md"
    report.write_text(
        "## 昨日选股复盘\n"
        "<table>\n"
        "<thead><tr><th>原排名</th><th>代码</th><th>股票</th><th>原评分</th><th>买入价</th>"
        "<th>最新价</th><th>持有N日</th><th>累计涨跌</th></tr></thead>\n"
        "<tbody>\n"
        "<tr><td>#6</td><td>600027</td><td>华电国际</td><td>86.08</td><td>5.62</td>"
        "<td>5.92</td><td>1</td><td>+5.34%</td></tr>\n"
        "</tbody></table>\n"
        "## 综合排名 TOP20\n"
        "<table>\n"
        "<thead><tr><th>排名</th><th>代码</th><th>股票名称</th><th>综合得分</th><th>详细分析</th></tr></thead>\n"
        "<tbody>\n"
        f"<tr><td>3</td><td>688111</td><td>金山办公</td><td>87.08</td><td>{_RICH_DETAIL}</td></tr>\n"
        "</tbody></table>\n",
        encoding="utf-8",
    )

    parsed = module._parse_opportunity_report(report)

    codes = [item["code"] for item in parsed["items"]]
    assert codes == ["688111"]  # 复盘行 600027 必须被排除
    item = parsed["items"][0]
    assert item["rating"] == "S"
    assert [field["label"] for field in item["fields"]][:2] == ["概览", "涨幅"]


def test_parse_opportunity_report_handles_outcome_marker_column(tmp_path, monkeypatch):
    """2026-07-29 起排名表在「综合得分」与「详细分析」之间插入了「表现标记」列。

    解析器不能再把第 5 格当作详细分析(否则整张排名表被 '【' 过滤器丢光,
    风险·机遇大屏的撮合矩阵/行业热力全空)。
    """
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = results_dir / "opportunity_top10_20260810_193543.md"
    report.write_text(
        "## 📅 昨日选股复盘\n"
        "<table>\n"
        "<thead><tr><th>原排名</th><th>代码</th><th>股票</th><th>原评分</th><th>买入价</th>"
        "<th>最新价</th><th>持有N日</th><th>累计涨跌</th></tr></thead>\n"
        "<tbody>\n"
        "<tr><td>#2</td><td>301239</td><td>普瑞眼科</td><td>89.39</td><td>33.32</td>"
        "<td>38.00</td><td>1</td><td>+14.05%</td></tr>\n"
        "</tbody></table>\n"
        "## 综合排名 TOP20\n"
        "<table>\n"
        "<thead><tr><th>排名</th><th>代码</th><th>股票名称</th><th>综合得分</th>"
        "<th>表现标记</th><th>详细分析</th></tr></thead>\n"
        "<tbody>\n"
        "<tr><td>1</td><td>688111</td><td>金山办公</td><td>87.08</td>"
        f"<td>偏弱体质📉技术乏力</td><td>{_RICH_DETAIL}</td></tr>\n"
        "</tbody></table>\n",
        encoding="utf-8",
    )

    parsed = module._parse_opportunity_report(report)

    codes = [item["code"] for item in parsed["items"]]
    assert codes == ["688111"]  # 复盘行 301239 仍须排除
    item = parsed["items"][0]
    assert item["rank"] == 1
    assert item["score"] == 87.08
    assert item["rating"] == "S"
    assert item["sector"] == "软件服务"
    assert [field["label"] for field in item["fields"]][:2] == ["概览", "涨幅"]


def test_parse_opportunity_report_prefers_market_env_line_over_review_heading(tmp_path, monkeypatch):
    """market_env should surface the 市场环境 blockquote, not the first 复盘 heading."""
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = results_dir / "opportunity_top10_20260602_095018.md"
    report.write_text(
        "## 📅 昨日选股复盘\n"
        "> **上一份报告**: (2026-05-31)\n"
        "## 综合排名 TOP20\n"
        "> **市场环境**: 🟡 中性 | 沪深300 近5日 -1.57% / 近20日 +0.70% | 今日 S 级 5 只\n"
        "<table>\n"
        "<tbody>\n"
        f"<tr><td>1</td><td>688111</td><td>金山办公</td><td>87.08</td><td>{_RICH_DETAIL}</td></tr>\n"
        "</tbody></table>\n",
        encoding="utf-8",
    )

    parsed = module._parse_opportunity_report(report)

    assert "市场环境" in parsed["market_env"]
    assert "中性" in parsed["market_env"]
    assert "昨日选股复盘" not in parsed["market_env"]


def test_parse_opportunity_report_extracts_run_meta(tmp_path, monkeypatch):
    """报告头的 kronos-run-meta 注释要解析为 run_meta 并透传 latest_report。"""
    module = _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = results_dir / "opportunity_top10_20260610_153102.md"
    report.write_text(
        '<!-- kronos-run-meta {"run_at": "2026-06-10T15:31:02", "source": "multi", '
        '"candidate_limit": 100, "ruleset_version": "v24", "config_hash": "abc123def456"} -->\n'
        "> 🧾 运行 2026-06-10 15:31 · 评分规则 v24 · 来源 multi\n\n"
        "## 综合排名 TOP20\n"
        "<table>\n"
        "<tbody>\n"
        f"<tr><td>1</td><td>688111</td><td>金山办公</td><td>87.08</td><td>{_RICH_DETAIL}</td></tr>\n"
        "</tbody></table>\n",
        encoding="utf-8",
    )

    parsed = module._parse_opportunity_report(report)
    assert parsed["run_meta"]["ruleset_version"] == "v24"
    assert parsed["run_meta"]["source"] == "multi"
    assert parsed["run_meta"]["config_hash"] == "abc123def456"

    cards = module.load_opportunity_report_cards(report.name)
    assert cards["latest_report"]["run_meta"]["ruleset_version"] == "v24"

    latest = module._load_latest_opportunities()
    assert latest["latest_report"]["run_meta"]["source"] == "multi"
    canvas = latest["canvas"]
    assert canvas["levels"] == ["报告", "板块", "股票", "分析内容"]
    assert any(view["id"] == "business_tag" for view in canvas["views"])
    assert canvas["stats"]["sectors"] == 1
    assert canvas["stats"]["stocks"] == 1
    assert {"root", "sector-1", "stock-1-688111"} <= {node["id"] for node in canvas["nodes"]}
    stock_node = next(node for node in canvas["nodes"] if node["id"] == "stock-1-688111")
    analysis_labels = [section["label"] for section in stock_node["analysis"]]
    # 概览/涨幅 were previously dropped by the [:10] cap; the canvas now surfaces
    # every 【…】 section in full (no truncation).
    assert analysis_labels[:3] == ["概览", "入选原因", "涨幅"]
    assert {"概览", "涨幅", "最新动态", "高级"} <= set(analysis_labels)
    assert len(analysis_labels) == 12


def test_opportunity_canvas_uses_full_run_sorted_and_hot_sector_memberships(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)
    from data_store.schema import migrate
    from data_store import hot_sector_repo, opportunity_repo

    conn = sqlite3.connect(tmp_path / "kronos_test.sqlite", isolation_level=None)
    conn.row_factory = sqlite3.Row
    migrate(conn)
    _patch_store_conn(monkeypatch, conn)

    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = _write_rich_report(results_dir, "opportunity_top10_20260614_101500.md")
    opportunity_repo.save_run(
        {
            "run_at": "2026-06-14T10:15:00",
            "run_date": "2026-06-14",
            "source": "multi",
            "candidate_limit": 100,
            "mode": "market_scan",
            "report_file": report.name,
        },
        [
            {"code": "600001", "name": "高分股", "total_score": 91.2, "rating": "S", "source": "heat"},
            {"code": "688111", "name": "金山办公", "total_score": 87.08, "rating": "S", "source": "multi"},
            {"code": "000001", "name": "平安银行", "total_score": 73.5, "rating": "B", "source": "moneyflow"},
        ],
    )
    hot_sector_repo.save_snapshot(
        boards=[
            {"code": "BK1001", "name": "软件服务", "type": "行业", "rank": 1, "main_net_inflow": 1.2e8},
            {"code": "BK2001", "name": "人工智能", "type": "概念", "rank": 2, "main_net_inflow": 8.8e7},
        ],
        stocks=[
            {"sector_code": "BK1001", "code": "688111", "name": "金山办公", "sector_stock_rank": 3, "main_net_inflow": 3e7},
            {"sector_code": "BK2001", "code": "688111", "name": "金山办公", "sector_stock_rank": 8, "main_net_inflow": 1.1e7},
            {"sector_code": "BK1001", "code": "600001", "name": "高分股", "sector_stock_rank": 1, "main_net_inflow": 6e7},
        ],
        relations=[
            {"board_code": "BK1001", "code": "688111", "relation_type": "dragon_tiger", "trade_date": "2026-06-14", "amount": 5e7},
        ],
        meta={"source": "sector_hot", "board_limit": 10},
    )

    payload = module._load_latest_opportunities()

    assert [item["code"] for item in payload["items"]] == ["600001", "688111", "000001"]
    assert payload["stats"]["total"] == 3
    canvas = payload["canvas"]
    assert any(view["id"] == "score_rank" for view in canvas["views"])
    stock_node = next(node for node in canvas["nodes"] if node.get("stock_code") == "688111" and node["type"] == "stock")
    assert stock_node["score_rank"] == 2
    assert stock_node["sector_rank"] == 2
    assert stock_node["sector_peer_count"] == 2
    assert stock_node["hot_sector_count"] == 2
    assert "软件服务 板块#1 / 成分股#3" in stock_node["hot_sector_rank_summary"]
    assert stock_node["lhb_hit"] is True
    assert {"from": stock_node["id"], "to": "hot-board-BK1001", "relation": "stock_hot_board"} in canvas["edges"]

    conn.close()


def test_export_opportunity_canvas_excel_without_hot_sector_snapshot(tmp_path, monkeypatch):
    module = _load_webui_core(tmp_path, monkeypatch)
    from data_store.schema import migrate
    from data_store import opportunity_repo

    conn = sqlite3.connect(tmp_path / "kronos_test.sqlite", isolation_level=None)
    conn.row_factory = sqlite3.Row
    migrate(conn)
    _patch_store_conn(monkeypatch, conn)

    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    report = _write_rich_report(results_dir, "opportunity_top10_20260614_111500.md")
    opportunity_repo.save_run(
        {
            "run_at": "2026-06-14T11:15:00",
            "run_date": "2026-06-14",
            "source": "multi",
            "candidate_limit": 100,
            "mode": "market_scan",
            "report_file": report.name,
        },
        [
            {"code": "688111", "name": "金山办公", "total_score": 87.08, "rating": "S", "source": "multi"},
            {"code": "000001", "name": "平安银行", "total_score": 73.5, "rating": "B", "source": "moneyflow"},
        ],
    )

    # Suite enrichment recomputes per-stock via STOCK_SUITE_SERVICE.get_suite, which would
    # hit the network. Stub it: full suite for 688111, degraded for anything else — so the
    # export exercises both the enriched and the baseline-fallback paths deterministically.
    def _fake_get_suite(code, name="", force_refresh=False):
        if code != "688111":
            return {"success": False, "error": "offline", "stock": {"code": code}}
        return {
            "success": True,
            "overview": {
                "radar": {"main_force_phase": {"score": 72, "label": "强势主导"},
                          "control_degree": {"score": 68, "label": "中控"}},
                "key_signals": [{"label": "仓位上限", "value": "60%"},
                                {"label": "大盘周期", "value": "震荡期"}],
                "scenario_probability": {"bullish": 60, "bearish": 13, "sideways": 27},
            },
            "risk_control": {
                "available": True,
                "execution_plan": {
                    "stop_loss": {"price": 28.5, "drop_pct": 7.5, "basis": "ATR(20)×1.5 下沿"},
                    "risk_reward": {"ratio": "1:2.6", "expected_return_pct": 19.5},
                },
                "scaled_entry": [{"label": "现价建仓", "price": 30.8, "position_pct": 10}],
                "tiered_take_profit": [{"label": "第一止盈(前高)", "price": 36.8, "sell_pct": 30}],
                "deep_signals": [{"text": "最大回撤：当前 -8.5%"},
                                 {"text": "波动率(60日年化) 42.3%"},
                                 {"text": "夏普比率(60日年化) 1.35"}],
            },
            "quant_matrix": {"data_status": "fresh", "current_posture": "震荡偏多",
                             "signals_matrix": [{"model": "海龟交易", "period": "daily", "signal": 1}]},
            "chip_control": {"data_status": "fresh", "control_degree": 68, "control_label": "中控",
                             "concentration_90": 12.0},
            "main_force_deep": {"dragon_tiger": {"quant_seat_appearances": 2, "net_inst_buy_30d": 5e7},
                                "hsgt": {"latest": {"hold_ratio": 3.5}}},
            "institutional_holdings": {"top10_floatholders": {"concentration": 42.0},
                                       "holder_number": {"latest_num": 32000}},
        }

    monkeypatch.setattr(module.STOCK_SUITE_SERVICE, "get_suite", _fake_get_suite)

    out = module.export_opportunity_canvas_excel()

    assert out["file"].startswith("opportunity_canvas_")
    import openpyxl
    wb = openpyxl.load_workbook(out["path"], read_only=True)
    assert {"报告信息", "股票排名", "股票全量明细", "板块汇总", "热门板块关联",
            "分析内容", "画布节点", "画布关系"} <= set(wb.sheetnames)
    # 6 张新 sheet 一定存在
    assert {"投资速览", "评分雷达", "风控执行计划", "量化模型矩阵", "筹码与机构",
            "术语表与图例"} <= set(wb.sheetnames)
    rank_header = list(wb["股票排名"].iter_rows(values_only=True))[0]
    assert rank_header[:4] == ("全量排名", "报告排名", "代码", "名称")
    # 细化列:涨幅拆分 + 评分分项 + 关键信号 + 来源细节 + 降级标记(即使本行无值也保留表头)
    for col in ("当日涨幅", "3日涨幅", "5日涨幅", "板块分", "技术分", "量化分",
                "追高风险", "RSI", "卖出信号", "量化总分", "来源细节", "是否降级"):
        assert col in rank_header
    rows = list(wb["股票排名"].iter_rows(values_only=True))
    assert rows[1][2] == "688111"
    # 全量明细:每个 【…】 字段单独成列 + 量化模型 + 分析摘要
    detail_header = list(wb["股票全量明细"].iter_rows(values_only=True))[0]
    for col in ("量化模型", "概览", "涨幅", "技术", "基本面", "高级", "历史重复入选", "分析摘要"):
        assert col in detail_header
    # 空的热门板块关联表也应带列头,而非 0x0
    assert list(wb["热门板块关联"].iter_rows(values_only=True))[0][0] == "代码"

    # 投资速览:688111 复算成功 → 带白话列与风控数值
    ov_header = list(wb["投资速览"].iter_rows(values_only=True))[0]
    for col in ("建议操作", "现价", "止损价", "第一目标价", "盈亏比", "夏普"):
        assert col in ov_header
    ov_rows = list(wb["投资速览"].iter_rows(values_only=True))[1:]
    ov_by_code = {r[ov_header.index("代码")]: r for r in ov_rows}
    assert "688111" in ov_by_code
    enriched = ov_by_code["688111"]
    assert enriched[ov_header.index("现价")] == 30.8
    assert enriched[ov_header.index("止损价")] == 28.5
    assert enriched[ov_header.index("盈亏比")] == "1:2.6"
    assert enriched[ov_header.index("建议操作")] in ("买入", "观望", "回避")

    # 报告信息:折算覆盖率字段
    info_header = list(wb["报告信息"].iter_rows(values_only=True))[0]
    for col in ("复算成功数", "复算失败数", "复算覆盖率"):
        assert col in info_header
    info_row = list(wb["报告信息"].iter_rows(values_only=True))[1]
    assert info_row[info_header.index("复算成功数")] >= 1

    # 术语表非空且含关键术语
    glossary = list(wb["术语表与图例"].iter_rows(values_only=True))
    terms = {r[0] for r in glossary[1:]}
    assert {"建议操作", "夏普比率", "盈亏比"} <= terms

    wb.close()
    conn.close()


def test_stock_dashboard_route_survives_existing_report(tmp_path, monkeypatch):
    _load_webui_core(tmp_path, monkeypatch)
    results_dir = tmp_path / "results"
    results_dir.mkdir(parents=True, exist_ok=True)
    (results_dir / "opportunity_top10_20260521_101600.md").write_text(
        "| 排名 | 代码 | 股票名称 | 综合得分 | 详细分析 |\n"
        "| --- | --- | --- | --- | --- |\n"
        "| 1 | 600519 | 贵州茅台 | 42.5 | 【概览】评级C，建议：观望；【入选原因】测试 |\n",
        encoding="utf-8",
    )

    from robyn.testing import TestClient

    sys.modules.pop("webui.robyn_app", None)
    robyn_module = importlib.import_module("webui.robyn_app")

    with TestClient(robyn_module.app) as client:
        response = client.get("/api/stock-dashboard")

    assert response.status_code == 200
    assert response.json()["opportunity"]["items"][0]["code"] == "600519"


def test_opportunity_report_route_is_wired_and_maps_missing_to_404(tmp_path, monkeypatch):
    # NOTE: robyn's TestClient does not strip the query string before route
    # matching, so `?file=` cannot be exercised here (it returns robyn's default
    # "Not Found"). The parsing/validation is covered by the unit tests above; the
    # `?file=` happy-path is smoke-tested over live HTTP. This test proves the route
    # is registered and that a missing report maps to our JSON 404 envelope.
    _load_webui_core(tmp_path, monkeypatch)

    from robyn.testing import TestClient

    sys.modules.pop("webui.robyn_app", None)
    robyn_module = importlib.import_module("webui.robyn_app")

    with TestClient(robyn_module.app) as client:
        response = client.get("/api/opportunity-report")

    assert response.status_code == 404
    assert response.json()["error"] == "Report not found"
